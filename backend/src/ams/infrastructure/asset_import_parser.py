"""Bounded CSV/XLSX parser for the baseline Asset import template."""

from __future__ import annotations

import csv
import io
from collections.abc import Mapping
from pathlib import PurePath
from uuid import UUID
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ams.application.asset_imports import (
    MAX_IMPORT_FILE_SIZE_BYTES,
    AssetImportFile,
    IAssetImportParser,
    UnsupportedImportFileError,
)
from ams.domain.asset_imports import AssetImportRow, BaselineAssetData
from ams.domain.assets import AssetCondition

MAX_IMPORT_ROWS = 1000
TEMPLATE_COLUMNS = (
    "name",
    "description",
    "source",
    "total_quantity",
    "category_id",
    "responsible_team_id",
    "custodian_user_id",
    "condition",
    "condition_note",
)
CSV_CONTENT_TYPES = frozenset({"text/csv", "application/csv", "application/vnd.ms-excel"})
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class CsvXlsxAssetImportParser(IAssetImportParser):
    """Parse the frozen nine-column template with bounded resource use."""

    def parse(self, source: AssetImportFile) -> tuple[AssetImportRow, ...]:
        """Validate file identity, template header, row limit, and scalar fields."""

        if not source.content or len(source.content) > MAX_IMPORT_FILE_SIZE_BYTES:
            raise UnsupportedImportFileError("Import file must be between 1 byte and 5 MB.")
        suffix = PurePath(source.filename).suffix.casefold()
        if suffix == ".csv" and source.content_type in CSV_CONTENT_TYPES:
            raw_rows = self._csv_rows(source.content)
        elif suffix == ".xlsx" and source.content_type == XLSX_CONTENT_TYPE:
            raw_rows = self._xlsx_rows(source.content)
        else:
            raise UnsupportedImportFileError("Upload a CSV or XLSX file.")
        return self._parse_rows(raw_rows)

    @staticmethod
    def _csv_rows(content: bytes) -> list[list[object]]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exception:
            raise UnsupportedImportFileError("CSV file must use UTF-8 encoding.") from exception
        try:
            return [list(row) for row in csv.reader(io.StringIO(text, newline=""))]
        except csv.Error as exception:
            raise UnsupportedImportFileError("CSV file is malformed.") from exception

    @staticmethod
    def _xlsx_rows(content: bytes) -> list[list[object]]:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            if sheet is None:
                raise UnsupportedImportFileError("XLSX file has no active worksheet.")
            rows: list[list[object]] = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(list(row))
            workbook.close()
            return rows
        except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exception:
            raise UnsupportedImportFileError("XLSX file is malformed.") from exception

    @staticmethod
    def _parse_rows(raw_rows: list[list[object]]) -> tuple[AssetImportRow, ...]:
        if not raw_rows:
            raise UnsupportedImportFileError("Import file has no template header.")
        headers = tuple(_cell_text(value).casefold() for value in raw_rows[0])
        if headers != TEMPLATE_COLUMNS:
            raise UnsupportedImportFileError(
                "Import header does not match the required template."
            )
        populated = [row for row in raw_rows[1:] if any(_cell_text(value) for value in row)]
        if not populated:
            raise UnsupportedImportFileError("Import file has no data rows.")
        if len(populated) > MAX_IMPORT_ROWS:
            raise UnsupportedImportFileError("Import file exceeds the 1000-row limit.")
        results: list[AssetImportRow] = []
        for row_number, values in enumerate(raw_rows[1:], start=2):
            if not any(_cell_text(value) for value in values):
                continue
            padded = values[: len(TEMPLATE_COLUMNS)] + [None] * max(
                0, len(TEMPLATE_COLUMNS) - len(values)
            )
            record = dict(zip(TEMPLATE_COLUMNS, padded, strict=True))
            results.append(_parse_record(row_number, record))
        return tuple(results)


def _parse_record(row_number: int, record: Mapping[str, object]) -> AssetImportRow:
    errors: list[str] = []
    name = _required_text(record["name"], "name", 200, errors)
    description = _optional_text(record["description"], "description", 1000, errors)
    source = _required_text(record["source"], "source", 255, errors)
    quantity = _positive_integer(record["total_quantity"], "total_quantity", errors)
    category_id = _required_uuid(record["category_id"], "category_id", errors)
    team_id = _optional_uuid(record["responsible_team_id"], "responsible_team_id", errors)
    custodian_id = _required_uuid(record["custodian_user_id"], "custodian_user_id", errors)
    condition = _condition(record["condition"], errors)
    condition_note = _optional_text(record["condition_note"], "condition_note", 1000, errors)
    if errors:
        return AssetImportRow(row_number, None, tuple(errors))
    assert name is not None
    assert source is not None
    assert quantity is not None
    assert category_id is not None
    assert custodian_id is not None
    assert condition is not None
    return AssetImportRow(
        row_number,
        BaselineAssetData(
            name,
            description,
            source,
            quantity,
            category_id,
            team_id,
            custodian_id,
            condition,
            condition_note,
        ),
        (),
    )


def _cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _required_text(value: object, field: str, maximum: int, errors: list[str]) -> str | None:
    result = _cell_text(value)
    if not result:
        errors.append(f"{field} is required")
        return None
    if len(result) > maximum:
        errors.append(f"{field} exceeds {maximum} characters")
        return None
    return result


def _optional_text(value: object, field: str, maximum: int, errors: list[str]) -> str | None:
    result = _cell_text(value)
    if not result:
        return None
    if len(result) > maximum:
        errors.append(f"{field} exceeds {maximum} characters")
        return None
    return result


def _positive_integer(value: object, field: str, errors: list[str]) -> int | None:
    text = _cell_text(value)
    try:
        number = int(text)
    except ValueError:
        errors.append(f"{field} must be a positive integer")
        return None
    if str(number) != text and not (isinstance(value, float) and value.is_integer()):
        errors.append(f"{field} must be a positive integer")
        return None
    if number <= 0:
        errors.append(f"{field} must be a positive integer")
        return None
    return number


def _required_uuid(value: object, field: str, errors: list[str]) -> UUID | None:
    text = _cell_text(value)
    if not text:
        errors.append(f"{field} is required")
        return None
    try:
        return UUID(text)
    except ValueError:
        errors.append(f"{field} must be a UUID")
        return None


def _optional_uuid(value: object, field: str, errors: list[str]) -> UUID | None:
    text = _cell_text(value)
    if not text:
        return None
    try:
        return UUID(text)
    except ValueError:
        errors.append(f"{field} must be a UUID")
        return None


def _condition(value: object, errors: list[str]) -> AssetCondition | None:
    text = _cell_text(value).upper()
    try:
        return AssetCondition(text or AssetCondition.GOOD.value)
    except ValueError:
        errors.append("condition must be GOOD, DAMAGED, or LOST")
        return None
